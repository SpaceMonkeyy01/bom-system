"""Seed the category tree, attribute definitions and the material master from Categorization.xlsx."""

from __future__ import annotations

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import AttributeDef, AttributeValue, Category, Material, MaterialExternalKey
from app.models.core import ExternalSystem
from app.seed.common import RAW_DIR, clean, ensure_uoms, norm_uom

CATEGORIZATION_XLSX = RAW_DIR / "Categorization.xlsx"

# attribute name -> (column index in Inventory Working File, data type, uom)
ATTRIBUTE_COLUMNS: dict[str, tuple[int, str, str | None]] = {
    "Grade": (30, "text", None),
    "Power Supply Type": (31, "text", None),
    "Electrical Rating": (32, "text", None),
    "Electrical Power": (33, "text", "W"),
    "Voltage": (34, "text", "V"),
    "Length": (35, "text", None),
    "Width": (37, "text", None),
    "Height": (39, "text", None),
    "Weight": (41, "number", "kg"),
    "Density": (42, "text", None),
    "Volume": (44, "text", None),
    "Sheet Size": (45, "text", None),
    "Gauge": (46, "text", None),
    "Thickness": (47, "text", None),
    "Color Rating": (48, "text", None),
    "Color": (49, "text", None),
    "Brand Company": (50, "text", None),
    "Finish": (51, "text", None),
    "Curing Speed": (52, "text", None),
    "Stain Color": (53, "text", None),
    "Source Code": (54, "text", None),
    "Brand Code": (55, "text", None),
    "Tint Color": (56, "text", None),
    "Roll Size": (57, "text", None),
    "Pack Size": (58, "text", None),
    "Source Company": (59, "text", None),
    "Shape": (61, "text", None),
    "Size/Number": (62, "text", None),
    "Specification": (63, "text", None),
    "HS Code": (24, "text", None),
    "Model Number / Brand Code": (25, "text", None),
    "Supplier Part Number": (26, "text", None),
    "Neon Type": (22, "text", None),
}


def seed_categories(db: Session, path=CATEGORIZATION_XLSX) -> dict[str, Category]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["CategoryMaster"]
    by_path: dict[str, Category] = {c.path: c for c in db.scalars(select(Category))}
    for row in ws.iter_rows(min_row=2, values_only=True):
        levels = [clean(x) for x in row[:4]]
        levels = [x for x in levels if x]
        if not levels:
            continue
        parent = None
        for depth, name in enumerate(levels, start=1):
            p = " / ".join(levels[:depth])
            cat = by_path.get(p)
            if not cat:
                cat = Category(parent=parent, name=name, level=depth, path=p)
                db.add(cat)
                db.flush()
                by_path[p] = cat
            parent = cat
    return by_path


def _norm_path(p: str | None) -> str | None:
    if not p:
        return None
    return " / ".join(part.strip() for part in str(p).split("/") if part.strip())


def seed_materials(db: Session, path=CATEGORIZATION_XLSX) -> dict:
    """Load Inventory Working File into Material + external keys + attributes. Idempotent on Odoo reference."""
    uoms = ensure_uoms(db)
    cats = seed_categories(db, path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Inventory Working File"]

    defs: dict[str, AttributeDef] = {d.name: d for d in db.scalars(select(AttributeDef).where(AttributeDef.category_id.is_(None)))}
    for i, (name, (_, dtype, uom)) in enumerate(ATTRIBUTE_COLUMNS.items()):
        if name not in defs:
            d = AttributeDef(category_id=None, name=name, data_type=dtype, uom=uom, sort_order=i)
            db.add(d)
            defs[name] = d
    db.flush()

    existing_by_odoo: dict[str, Material] = {}
    for k in db.scalars(select(MaterialExternalKey).where(MaterialExternalKey.system == ExternalSystem.odoo)):
        existing_by_odoo[k.external_id] = k.material

    stats = {"created": 0, "updated": 0, "skipped": 0, "no_category": 0, "keys": 0, "attrs": 0, "duplicate_iflow": []}
    claimed: dict[tuple, int] = {
        (k.system, k.store, k.external_id): k.material_id for k in db.scalars(select(MaterialExternalKey))
    }
    rows = ws.iter_rows(min_row=3, values_only=True)
    for row in rows:
        name = clean(row[0])
        odoo_ref = row[2]
        if not name or odoo_ref in (None, ""):
            stats["skipped"] += 1
            continue
        odoo_ref = str(int(odoo_ref)) if isinstance(odoo_ref, float) else str(odoo_ref).strip()
        action = clean(row[1]) or ""
        status = clean(row[10]) or ""
        cat_path = _norm_path(clean(row[15]))
        category = cats.get(cat_path) if cat_path else None
        if cat_path and not category:
            # path not in master (e.g. "In Active (No Stock & Consumption)" variants) -> try level 1 only
            l1 = _norm_path(clean(row[11]))
            category = cats.get(l1) if l1 else None
        if not category:
            stats["no_category"] += 1
        inactive = action == "Not Needed" or status.lower().startswith("in active") or (cat_path or "").startswith("In Active")
        review = {"Done": "done", "To Review": "to_review", "Not Needed": "not_needed"}.get(action, "to_review")
        uom_code = norm_uom(clean(row[17]) or clean(row[8]))

        mat = existing_by_odoo.get(odoo_ref)
        if mat is None:
            mat = Material(name=name)
            db.add(mat)
            stats["created"] += 1
        else:
            stats["updated"] += 1
        mat.name = name
        mat.category = category
        mat.uom = uoms[uom_code]
        mat.product_type = clean(row[16]) or "Goods"
        mat.country_of_origin = clean(row[19])
        mat.supplier = clean(row[20])
        mat.supplier_city = clean(row[21])
        mat.brand = clean(row[23])
        mat.active = not inactive
        mat.review_status = review
        db.flush()

        have = {(k.system, k.store): k for k in mat.external_keys}

        def put_key(system: ExternalSystem, store: str, ext_id, ext_name=None):
            if ext_id in (None, "", "#N/A", "-"):
                return
            if isinstance(ext_id, float):
                ext_id = str(int(ext_id))
            ext_id = str(ext_id).strip()
            if not ext_id or ext_id.upper().startswith("#"):
                return
            owner = claimed.get((system, store, ext_id))
            if owner is not None and owner != mat.id:
                # Two Odoo products point at the same iFlow item (Odoo duplicates). First one wins; flag the rest.
                stats["duplicate_iflow"].append((name, odoo_ref, system.value, store, ext_id))
                mat.review_status = "to_review"
                mat.notes = f"Duplicate: {system.value} {store} {ext_id} already belongs to material #{owner}"
                return
            k = have.get((system, store))
            if k:
                k.external_id = ext_id
                k.external_name = ext_name
            else:
                k = MaterialExternalKey(material=mat, system=system, store=store, external_id=ext_id, external_name=ext_name)
                db.add(k)
                have[(system, store)] = k
                stats["keys"] += 1
            claimed[(system, store, ext_id)] = mat.id

        put_key(ExternalSystem.odoo, "", odoo_ref, name)
        old_cat = clean(row[6]) or ""
        # E = iFlow code in the BS store, F = in the NS store. D is the generic code for whichever store the item lives in.
        bs_code, ns_code, generic = row[4], row[5], row[3]
        if bs_code not in (None, "", "#N/A", "-"):
            put_key(ExternalSystem.iflow, "BS", bs_code, name)
        elif old_cat == "Business Signs" and generic not in (None, "", "#N/A", "-"):
            put_key(ExternalSystem.iflow, "BS", generic, name)
        if ns_code not in (None, "", "#N/A", "-"):
            put_key(ExternalSystem.iflow, "NS", ns_code, name)
        elif old_cat == "Neon Signs" and generic not in (None, "", "#N/A", "-"):
            put_key(ExternalSystem.iflow, "NS", generic, name)

        have_attrs = {a.def_id: a for a in mat.attributes}
        for attr_name, (col, _, _) in ATTRIBUTE_COLUMNS.items():
            val = clean(row[col]) if col < len(row) else None
            if not val:
                continue
            d = defs[attr_name]
            a = have_attrs.get(d.id)
            if a:
                a.value = val[:255]
            else:
                db.add(AttributeValue(material=mat, definition=d, value=val[:255]))
                stats["attrs"] += 1
    db.flush()
    return stats


def seed_suppliers(db: Session) -> dict:
    """Create Supplier rows from the free-text supplier names on materials and link them as preferred supplier."""
    from app.models import Supplier

    existing = {s.name.lower(): s for s in db.scalars(select(Supplier))}
    created = linked = 0
    for m in db.scalars(select(Material).where(Material.supplier.isnot(None))):
        for raw in str(m.supplier).split(","):
            name = clean(raw)
            if not name or len(name) < 2:
                continue
            s = existing.get(name.lower())
            if not s:
                s = Supplier(name=name, city=m.supplier_city)
                db.add(s)
                db.flush()
                existing[name.lower()] = s
                created += 1
            if m.preferred_supplier_id is None:
                m.preferred_supplier_id = s.id
                linked += 1
    db.flush()
    return {"suppliers_created": created, "materials_linked": linked}
