"""Seed sign types, the historical-name mapping and template versions from BOM_Check.xlsx."""

from __future__ import annotations

import re
from datetime import datetime, timezone

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.engine.formula import check_formula, is_cost_formula
from app.models import ComponentLine, Material, MaterialExternalKey, Parameter, SignType, Template
from app.models.core import ExternalSystem
from app.seed.common import RAW_DIR, clean, cost_tag_for_category, norm_uom

BOM_CHECK_XLSX = RAW_DIR / "BOM_Check.xlsx"
NON_TEMPLATE_SHEETS = {"Work TypeCenter", "Parameters", "Sign Types", "Parameter Ali Hassan", "Inventory data", "Sheet28"}
SECOND_SERIES_MARKER = "Parameter Ali Hassan"  # sheets after this are the second authoring pass

ORDER_PARAMS = ["Length_in", "Width_in", "Area_sqft"]


def _cell(txt: list[str], i: int) -> str:
    return txt[i] if 0 <= i < len(txt) else ""


def sanitize_param(name: str) -> str:
    """Formulas need identifiers: 'Alu/SS face Crunched Area' -> 'Alu_SS_face_Crunched_Area'."""
    s = re.sub(r"[^A-Za-z0-9_]+", "_", name.strip()).strip("_")
    if s and s[0].isdigit():
        s = "_" + s
    return s or "Param"


def rewrite_formula(formula: str, renames: dict[str, str]) -> str:
    """Replace raw (non-identifier) parameter names in a formula with their sanitized form, longest first."""
    out = formula
    for raw, new in sorted(renames.items(), key=lambda kv: -len(kv[0])):
        if raw != new and raw in out:
            out = out.replace(raw, new)
    return out


def _param_name_from_label(label: str) -> str:
    s = re.sub(r"\(.*?\)", "", label)
    s = re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_")
    return s or "Param"


def parse_template_sheet(ws) -> dict | None:
    """Return {name, params:[...], lines:[...]} for one template sheet, or None if it holds nothing."""
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    name = None
    params: list[dict] = []
    lines: list[dict] = []
    mode = None
    hdr: dict[str, int] = {}
    seen_component_header = 0
    for r in rows:
        txt = [str(c).strip() if c is not None else "" for c in r]
        if not name and "BOM Template" in txt:
            name = _cell(txt, txt.index("BOM Template") + 1) or None
        if "Field Name" in txt and "Parameter Name" in txt and not params:
            mode = "params"
            hdr = {h: i for i, h in enumerate(txt) if h}
            continue
        if "Component List" in txt and "Formula" in txt:
            seen_component_header += 1
            if seen_component_header > 1:  # the Box Making section has its own component list; not a BOM line
                mode = None
                break
            mode = "lines"
            hdr = {h: i for i, h in enumerate(txt) if h}
            continue
        if mode == "params":
            fn, pn = _cell(txt, hdr["Field Name"]), _cell(txt, hdr["Parameter Name"])
            if not fn and not pn:
                if params:
                    mode = None
                continue
            itype = _cell(txt, hdr.get("Input Type", -1)).lower()
            default = _cell(txt, hdr.get("Input Value", -1))
            is_select = "dropdown" in itype or fn.lower() == "finish"
            params.append(
                {
                    "label": fn or pn,
                    "param_name": (pn or _param_name_from_label(fn)).strip(),
                    "uom": _cell(txt, hdr.get("UOM", -1)) or None,
                    "input_type": "select" if is_select else "number",
                    "options": ["Matt", "Gloss", "Satin"] if is_select and fn.lower() == "finish" else [],
                    "default": default or None,
                }
            )
        elif mode == "lines":
            first = txt[0] if txt else ""
            if first and not re.match(r"^\d+(\.0)?$", first):
                mode = None
                continue
            if not first:
                continue

            def g(prefix: str) -> str:
                for k, i in hdr.items():
                    if k.lower().startswith(prefix.lower()):
                        return _cell(txt, i)
                return ""

            sku = g("Product/SKU ID")
            sku = str(int(float(sku))) if re.match(r"^\d+(\.0)?$", sku) else (sku or None)
            uc = g("Unit Cost")
            qty_cell = g("Quantity")
            formula = g("Formula") or None
            if not formula and re.match(r"^-?\d+(\.\d+)?$", qty_cell or ""):
                formula = qty_cell  # a constant typed straight into the Quantity cell (e.g. 1 power supply)
            lines.append(
                {
                    "odoo_ref": sku,
                    "formula": formula,
                    "usage": g("Component Usage") or g("Usage") or None,
                    "cost_tag": g("Product Cost Tag") or g("Cost Tag") or None,
                    "unit_cost": uc or None,
                    "work_centre": g("Work Type") or None,
                    "uom_hint": g("Odoo UOM") or g("UOM") or None,
                }
            )
    if not name and not lines:
        return None
    return {"name": name or ws.title, "params": params, "lines": lines}


def _sign_types_mapping(wb) -> tuple[dict[str, str], list[str]]:
    """From the 'Sign Types' tab: historical name -> canonical template name, and the canonical list."""
    ws = wb["Sign Types"]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    canonical: list[str] = []
    for r in rows[2:]:
        a = clean(r[0]) if r else None
        if a and a not in ("To Add Later",):
            canonical.append(a)
        if a == "To Add Later":
            break
    same_groups: dict[str, str] = {}  # "Same 1" -> canonical name (column M when present)
    mapping: dict[str, str] = {}
    for r in rows[1:]:
        hist = clean(r[5]) if len(r) > 5 else None
        if not hist:
            continue
        grp = clean(r[6]) if len(r) > 6 else None
        tmpl = clean(r[12]) if len(r) > 12 else None
        if grp and tmpl:
            same_groups[grp] = tmpl
    for r in rows[1:]:
        hist = clean(r[5]) if len(r) > 5 else None
        if not hist:
            continue
        grp = clean(r[6]) if len(r) > 6 else None
        tmpl = clean(r[12]) if len(r) > 12 else None
        target = tmpl or (same_groups.get(grp) if grp else None) or hist
        mapping[hist] = target
    return mapping, canonical


def _key(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", s.lower())


PAINT_GRAM_FORMULA = re.compile(r"40\.49|\*\s*22\s*\*")  # g/ft² coverage constants from the paint recipes


def _line_uom(ln: dict, mat, formula: str) -> str | None:
    """Line unit. The paint recipes compute grams even though Odoo stocks paint in quarts/gallons/sets,
    so those lines are Gram; everything else uses the sheet's UOM column or the material's unit."""
    if formula and PAINT_GRAM_FORMULA.search(formula):
        return "Gram"
    hint = ln.get("uom_hint")
    if not hint or str(hint).startswith("="):
        return mat.uom.code if mat and mat.uom else None
    return norm_uom(hint)


def seed_sign_types_and_templates(db: Session, path=BOM_CHECK_XLSX) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    mapping, canonical = _sign_types_mapping(wb)

    sign_types: dict[str, SignType] = {_key(s.name): s for s in db.scalars(select(SignType))}

    def get_sign_type(name: str, code: str | None = None) -> SignType:
        k = _key(name)
        st = sign_types.get(k)
        if not st:
            for s in sign_types.values():
                if any(_key(a) == k for a in s.aliases):
                    return s
            st = SignType(name=name, code=code, aliases=[])
            db.add(st)
            db.flush()
            sign_types[k] = st
        elif code and not st.code:
            st.code = code
        return st

    odoo_to_material: dict[str, Material] = {
        k.external_id: k.material for k in db.scalars(select(MaterialExternalKey).where(MaterialExternalKey.system == ExternalSystem.odoo))
    }

    stats = {"sign_types": 0, "templates": 0, "lines": 0, "unresolved_materials": [], "incomplete_templates": [], "aliases": 0}
    series = 1
    for ws in wb.worksheets:
        if ws.title == SECOND_SERIES_MARKER:
            series = 2
            continue
        if ws.title in NON_TEMPLATE_SHEETS:
            continue
        parsed = parse_template_sheet(ws)
        if not parsed or not parsed["lines"]:
            continue
        code = (re.match(r"^\s*(\d+)", ws.title) or [None, None])[1]
        st = get_sign_type(parsed["name"], code if series == 1 else None)
        source = f"BOM_Check.xlsx / {ws.title}"
        existing = db.scalar(select(Template).where(Template.sign_type_id == st.id, Template.source == source))
        if existing:
            continue  # idempotent: a sheet is imported once
        version = (max((t.version for t in st.templates), default=0) + 1) if st.templates else 1
        tmpl = Template(sign_type=st, version=version, status="draft", source=source, notes=f"Imported from {ws.title} (series {series})")
        db.add(tmpl)
        db.flush()
        stats["templates"] += 1

        seen_params: set[str] = set()
        renames = {p["param_name"]: sanitize_param(p["param_name"]) for p in parsed["params"]}
        for i, p in enumerate(parsed["params"]):
            pn = renames[p["param_name"]]
            if pn.lower() in seen_params:
                continue
            seen_params.add(pn.lower())
            db.add(
                Parameter(
                    template=tmpl,
                    seq=i,
                    label=p["label"],
                    param_name=pn,
                    uom=p["uom"],
                    input_type=p["input_type"],
                    options=p["options"],
                    default_value=p["default"],
                    default_from_order="finish" if pn.lower() == "finish" else None,
                    required=p["input_type"] == "number",
                )
            )
        param_names = list(renames.values()) + ORDER_PARAMS

        complete = True
        for i, ln in enumerate(parsed["lines"]):
            mat = odoo_to_material.get(ln["odoo_ref"]) if ln["odoo_ref"] else None
            if ln["odoo_ref"] and not mat:
                stats["unresolved_materials"].append((ws.title, ln["odoo_ref"]))
            formula = rewrite_formula(ln["formula"], renames) if ln["formula"] else None
            problem = None
            if not formula:
                problem = "no formula in source"
                formula = "0"
            else:
                problem = check_formula(formula, param_names)
            if not mat or problem:
                complete = False
            cat_path = mat.category.path if mat and mat.category else None
            cost_tag = ln["cost_tag"] or cost_tag_for_category(cat_path, mat.name if mat else None)
            uc = ln["unit_cost"]
            db.add(
                ComponentLine(
                    template=tmpl,
                    seq=i,
                    material=mat,
                    material_name_hint=(mat.name if mat else (f"Odoo {ln['odoo_ref']}" if ln["odoo_ref"] else "(no material)")),
                    qty_formula=formula,
                    unit_cost_formula=uc if (uc and is_cost_formula(uc)) else None,
                    uom=_line_uom(ln, mat, formula),
                    usage=ln["usage"],
                    cost_tag=cost_tag,
                    work_centre=ln["work_centre"],
                    optional=bool(problem),
                    notes=problem,
                )
            )
            stats["lines"] += 1
        if not complete:
            stats["incomplete_templates"].append(ws.title)
        elif series == 1 and not any(t.status == "published" for t in st.templates if t is not tmpl):
            tmpl.status = "published"
            tmpl.published_at = datetime.now(timezone.utc)
    db.flush()

    # historical Airtable names -> aliases on the canonical sign type
    for hist, target in mapping.items():
        st = get_sign_type(target)
        if _key(hist) != _key(st.name) and not any(_key(a) == _key(hist) for a in st.aliases):
            st.aliases = [*st.aliases, hist]
            stats["aliases"] += 1
    db.flush()
    stats["sign_types"] = len(sign_types)
    return stats
